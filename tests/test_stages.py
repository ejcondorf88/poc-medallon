import json
from pathlib import Path

import pandas as pd

import pytest
from openpyxl import load_workbook

from andinos_etl.artefactos import leer_plata
from andinos_etl.integridad import ErrorIntegridadArtifact, generar_checksums
from andinos_etl.stages import APPROVED_BUNDLE_NAME, ejecutar_bronce, ejecutar_oro, ejecutar_plata, verificar_artifact


def _crear_bronze(raiz: Path) -> Path:
    bronze = raiz / "bronce"
    bronze.mkdir(parents=True)
    tablas = {
        "tiendas": pd.DataFrame([{
            "tienda_id": "T001", "nombre_tienda": "Sucursal Central", "ciudad": "Quito",
            "provincia": "Pichincha", "formato": "Urbano", "fecha_apertura": "2020-01-01",
            "metros2": "1000",
        }]),
        "productos": pd.DataFrame([{
            "producto_id": "P001", "sku": "SKU001", "producto": "Producto", "categoria": "A",
            "subcategoria": "B", "marca": "Marca", "proveedor": "Proveedor",
            "costo_base": "2.00", "precio_lista": "5.00",
        }]),
        "clientes": pd.DataFrame([{
            "cliente_id": "C001", "segmento": "Premium", "ciudad_residencia": "Quito",
            "anio_nacimiento": "1990", "fecha_alta": "2021-01-01", "canal_preferido": "Web",
        }]),
        "promociones": pd.DataFrame([{
            "promocion_id": "PR001", "nombre_promocion": "Campaña", "fecha_inicio": "2025-01-01",
            "fecha_fin": "2026-12-31", "tipo_promocion": "Porcentaje", "descuento_objetivo": "0.10",
            "categoria_objetivo": "A",
        }]),
        "ventas": pd.DataFrame([
            {
                "venta_id": "V001", "fecha_hora": "2025-01-15 10:00:00", "tienda_id": "T001",
                "cliente_id": "C001", "producto_id": "P001", "cantidad": "2",
                "precio_unitario": "5.00", "descuento_pct": "0.00", "costo_unitario": "2.00",
                "promocion_id": "PR001", "medio_pago": "Tarjeta", "tipo_movimiento": "Venta",
            },
            {
                "venta_id": "V002", "fecha_hora": "2026-01-15 11:00:00", "tienda_id": "T001",
                "cliente_id": "C001", "producto_id": "P001", "cantidad": "3",
                "precio_unitario": "5.00", "descuento_pct": "0.00", "costo_unitario": "2.00",
                "promocion_id": "PR001", "medio_pago": "Efectivo", "tipo_movimiento": "Venta",
            },
        ]),
    }
    for tabla, df in tablas.items():
        df.to_csv(bronze / f"{tabla}.csv", index=False, encoding="utf-8-sig")
    return bronze


def _bundle_aprobado(raiz: Path) -> Path:
    bronze = _crear_bronze(raiz)
    bronze_artifact = raiz / "bronze-artifact"
    plata_artifact = raiz / "plata-artifact"
    output = raiz / "output"
    ejecutar_bronce(bronze, bronze_artifact)
    ejecutar_plata(bronze_artifact, plata_artifact)
    ejecutar_oro(bronze_artifact, plata_artifact, output)
    return output / APPROVED_BUNDLE_NAME


def test_round_trip_bronze_plata_conserva_tipos_y_trazabilidad(tmp_path: Path) -> None:
    bronze = _crear_bronze(tmp_path)
    bronze_artifact = tmp_path / "bronze-artifact"
    plata_artifact = tmp_path / "plata-artifact"

    resultado_bronce = ejecutar_bronce(bronze, bronze_artifact)
    resultado_plata = ejecutar_plata(bronze_artifact, plata_artifact)
    plata = leer_plata(plata_artifact)

    assert resultado_bronce.status == "OK"
    assert resultado_plata.status == "OK"
    assert verificar_artifact(bronze_artifact).ok
    assert verificar_artifact(plata_artifact).ok
    assert plata["ventas"].limpio["_fila_origen"].tolist() == [2, 3]
    assert plata["ventas"].limpio["fecha_hora"].dtype.kind == "M"
    assert plata["ventas"].limpio["fecha"].dtype == object
    assert plata["ventas"].limpio["cantidad"].dtype == "int64"
    assert plata["ventas"].limpio["precio_unitario"].dtype == "float64"
    assert plata["tiendas"].limpio["fecha_apertura"].dtype == object
    assert plata["clientes"].limpio["anio_nacimiento"].dtype == "Int64"


def test_plata_preserva_desconocido_legitimo_tras_round_trip(tmp_path: Path) -> None:
    bronze = _crear_bronze(tmp_path)
    ventas = pd.read_csv(bronze / "ventas.csv", encoding="utf-8-sig", dtype=str)
    ventas.loc[1, "cliente_id"] = ""
    ventas.to_csv(bronze / "ventas.csv", index=False, encoding="utf-8-sig")
    bronze_artifact = tmp_path / "bronze-artifact"
    plata_artifact = tmp_path / "plata-artifact"
    ejecutar_bronce(bronze, bronze_artifact)
    ejecutar_plata(bronze_artifact, plata_artifact)

    plata = leer_plata(plata_artifact)

    assert plata["ventas"].limpio["cliente_id"].isna().tolist() == [False, True]


def test_plata_detecta_archivo_inesperado(tmp_path: Path) -> None:
    bronze = _crear_bronze(tmp_path)
    bronze_artifact = tmp_path / "bronze-artifact"
    plata_artifact = tmp_path / "plata-artifact"
    ejecutar_bronce(bronze, bronze_artifact)
    ejecutar_plata(bronze_artifact, plata_artifact)
    (plata_artifact / "inesperado.csv").write_text("id\n1\n", encoding="utf-8")

    with pytest.raises(ErrorIntegridadArtifact, match="inesperados"):
        leer_plata(plata_artifact)


def test_oro_approved_crea_bundle_exacto(tmp_path: Path) -> None:
    bronze = _crear_bronze(tmp_path)
    bronze_artifact = tmp_path / "bronze-artifact"
    plata_artifact = tmp_path / "plata-artifact"
    output = tmp_path / "output"
    ejecutar_bronce(bronze, bronze_artifact)
    ejecutar_plata(bronze_artifact, plata_artifact)

    resultado = ejecutar_oro(bronze_artifact, plata_artifact, output)

    bundle = output / APPROVED_BUNDLE_NAME
    esperados = {
        "modelo_estrella.xlsx", "manifest.json", "checksums.sha256", "gate.json",
        "modelo_estrella/dim_fecha.csv", "modelo_estrella/dim_tienda.csv",
        "modelo_estrella/dim_producto.csv", "modelo_estrella/dim_cliente.csv",
        "modelo_estrella/dim_promocion.csv", "modelo_estrella/fact_ventas.csv",
        "calidad/reporte_calidad.csv", "calidad/reglas_limpieza.csv",
    }
    assert resultado.status == "APPROVED"
    assert resultado.gate is not None
    assert resultado.gate["status"] == "APPROVED"
    assert len(resultado.gate["controls"]) == 19
    assert resultado.manifest is not None
    assert resultado.manifest["rows_processed"] == 2
    assert resultado.manifest["fact_rows"] == 2
    assert set(resultado.manifest["rows_by_table"]) == {
        "dim_tienda", "dim_producto", "dim_cliente", "dim_promocion", "dim_fecha", "fact_ventas",
    }
    assert resultado.manifest["rows_by_table"]["fact_ventas"] == 2
    assert resultado.manifest["ventas_netas"] == 25.0
    assert resultado.manifest["costo_total"] == 10.0
    assert resultado.manifest["margen"] == 15.0
    assert resultado.manifest["git_commit"]
    assert resultado.manifest["created_at_utc"].endswith("Z")
    contenido_fact = (resultado.output_dir / "modelo_estrella" / "fact_ventas.csv").read_text(encoding="utf-8-sig")
    assert "T" in contenido_fact.splitlines()[1].split(",")[1]
    assert {ruta.relative_to(bundle).as_posix() for ruta in bundle.rglob("*") if ruta.is_file()} == esperados
    assert verificar_artifact(bundle).ok


def test_fallo_no_crea_bundle_approved(tmp_path: Path) -> None:
    bronze = _crear_bronze(tmp_path)
    bronze_artifact = tmp_path / "bronze-artifact"
    plata_artifact = tmp_path / "plata-artifact"
    output = tmp_path / "output"
    ejecutar_bronce(bronze, bronze_artifact)
    ejecutar_plata(bronze_artifact, plata_artifact)
    ventas = pd.read_csv(bronze / "ventas.csv", encoding="utf-8-sig", dtype=str)
    ventas.loc[len(ventas)] = ventas.iloc[0].to_dict()
    ventas.loc[len(ventas) - 1, "venta_id"] = "V999"
    ventas.to_csv(bronze / "ventas.csv", index=False, encoding="utf-8-sig")
    bronze_inconsistente = tmp_path / "bronce-inconsistente"
    ejecutar_bronce(bronze, bronze_inconsistente)

    resultado = ejecutar_oro(bronze_inconsistente, plata_artifact, output)

    assert resultado.status == "REJECTED"
    assert not (output / APPROVED_BUNDLE_NAME).exists()
    assert (output / "diagnostico" / "gate.json").is_file()
    assert (output / "diagnostico" / "calidad" / "reporte_calidad.csv").is_file()


def test_verificar_detecta_divergencia_entre_csv_y_excel(tmp_path: Path) -> None:
    bundle = _bundle_aprobado(tmp_path)
    libro = load_workbook(bundle / "modelo_estrella.xlsx")
    hoja = libro["fact_ventas"]
    hoja.append([0] * hoja.max_column)
    libro.save(bundle / "modelo_estrella.xlsx")
    generar_checksums(bundle)

    resultado = verificar_artifact(bundle)

    assert not resultado.ok
    assert any("Conteos distintos en fact_ventas" in error for error in resultado.errors)


def test_verificar_detecta_manifest_conteo_distinto_del_bundle(tmp_path: Path) -> None:
    bundle = _bundle_aprobado(tmp_path)
    ruta_manifest = bundle / "manifest.json"
    manifest = json.loads(ruta_manifest.read_text(encoding="utf-8"))
    manifest["rows_by_table"]["fact_ventas"] = 999
    ruta_manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    generar_checksums(bundle)

    resultado = verificar_artifact(bundle)

    assert not resultado.ok
    assert any(
        "Conteos distintos en fact_ventas: manifest 999 | CSV 2 | Excel 2" in error
        for error in resultado.errors
    )


def test_verificar_exige_rows_by_table_en_el_manifest(tmp_path: Path) -> None:
    bundle = _bundle_aprobado(tmp_path)
    ruta_manifest = bundle / "manifest.json"
    manifest = json.loads(ruta_manifest.read_text(encoding="utf-8"))
    manifest.pop("rows_by_table")
    ruta_manifest.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    generar_checksums(bundle)

    resultado = verificar_artifact(bundle)

    assert not resultado.ok
    assert any("rows_by_table" in error for error in resultado.errors)
