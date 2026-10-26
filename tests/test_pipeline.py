import threading
import time
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from andinos_etl.config import Settings
from andinos_etl.excepciones import ErrorConcurrencia
from andinos_etl.pipeline import PipelineMedallon


def crear_bronze(data_dir: Path) -> None:
    bronze = data_dir / "bronce"
    bronze.mkdir(parents=True)
    tablas = {
        "tiendas": pd.DataFrame([{
            "tienda_id": "T001", "nombre_tienda": "Sucursal Central", "ciudad": "Quito",
            "provincia": "Pichincha", "formato": "Urbano", "fecha_apertura": "2020-01-01",
            "metros2": "1000",
        }]),
        "productos": pd.DataFrame([{
            "producto_id": "P001", "sku": "SKU001", "producto": "Producto", "categoria": "A",
            "subcategoria": "B", "marca": "Marca", "proveedor": "Proveedor",
            "costo_base": "2.00", "precio_lista": "5.00",
        }]),
        "clientes": pd.DataFrame([{
            "cliente_id": "C001", "segmento": "Premium", "ciudad_residencia": "Quito",
            "anio_nacimiento": "1990", "fecha_alta": "2021-01-01", "canal_preferido": "Web",
        }]),
        "promociones": pd.DataFrame([{
            "promocion_id": "PR001", "nombre_promocion": "Campaña", "fecha_inicio": "2025-01-01",
            "fecha_fin": "2026-12-31", "tipo_promocion": "Porcentaje", "descuento_objetivo": "0.10",
            "categoria_objetivo": "A",
        }]),
        "ventas": pd.DataFrame([
            {
                "venta_id": "V001", "fecha_hora": "2025-01-15 10:00:00", "tienda_id": "T001",
                "cliente_id": "C001", "producto_id": "P001", "cantidad": "2",
                "precio_unitario": "5.00", "descuento_pct": "0.00", "costo_unitario": "2.00",
                "promocion_id": "PR001", "medio_pago": "Tarjeta", "tipo_movimiento": "Venta",
            },
            {
                "venta_id": "V002", "fecha_hora": "2026-01-15 11:00:00", "tienda_id": "T001",
                "cliente_id": "C001", "producto_id": "P001", "cantidad": "3",
                "precio_unitario": "5.00", "descuento_pct": "0.00", "costo_unitario": "2.00",
                "promocion_id": "PR001", "medio_pago": "Efectivo", "tipo_movimiento": "Venta",
            },
        ]),
    }
    for tabla, df in tablas.items():
        df.to_csv(bronze / f"{tabla}.csv", index=False, encoding="utf-8-sig")


def _snapshot_publicado(data_dir: Path) -> dict[str, object]:
    snapshot: dict[str, object] = {}
    for capa in ("plata", "oro", "calidad"):
        for ruta in sorted((data_dir / capa).rglob("*.csv")):
            snapshot[str(ruta.relative_to(data_dir))] = ruta.read_bytes()
    libro = load_workbook(data_dir / "oro" / "modelo_estrella.xlsx", read_only=True, data_only=True)
    for hoja in libro.sheetnames:
        snapshot[f"xlsx:{hoja}"] = tuple(tuple(fila) for fila in libro[hoja].iter_rows(values_only=True))
    libro.close()
    return snapshot


def test_ejecucion_es_idempotente(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path / "data")
    crear_bronze(settings.data_dir)
    pipeline = PipelineMedallon(settings)


    pipeline.ejecutar()
    primera = _snapshot_publicado(settings.data_dir)
    pipeline.ejecutar()
    segunda = _snapshot_publicado(settings.data_dir)

    assert segunda == primera
    assert not settings.staging_dir.exists()


def test_salida_explica_capas_reglas_y_controles(tmp_path: Path, capsys) -> None:
    settings = Settings(data_dir=tmp_path / "data")
    crear_bronze(settings.data_dir)

    PipelineMedallon(settings).ejecutar()

    salida = capsys.readouterr().out
    assert "[BRONCE] Lee la fuente sin modificarla." in salida
    assert "[PLATA] Limpia cada tabla sin borrar evidencia:" in salida
    assert "[PLATA]   rechazar:signo_cantidad_incoherente:" in salida
    assert "[ORO] Paso 2/4:" in salida
    assert "[CALIDAD] C1 Conciliación de filas |" in salida
    assert "qué protege: Filas de bronce = filas de plata + rechazadas" in salida
    assert "cualquier FALLO bloquea la publicación." in salida


def test_ejecucion_concurrente_queda_serializada(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path / "data")
    crear_bronze(settings.data_dir)
    pipeline = PipelineMedallon(settings, lock_timeout=3.0, lock_poll=0.01)

    primera_entrada = threading.Event()
    liberar = threading.Event()
    entradas = []

    def ejecucion_simulada() -> None:
        entradas.append(threading.current_thread().name)
        primera_entrada.set()
        assert liberar.wait(3.0)

    pipeline._ejecutar_bloqueado = ejecucion_simulada
    hilos = [
        threading.Thread(target=pipeline.ejecutar, name="primera"),
        threading.Thread(target=pipeline.ejecutar, name="segunda"),
    ]
    hilos[0].start()
    assert primera_entrada.wait(1.0)
    hilos[1].start()
    time.sleep(0.15)

    assert len(entradas) == 1
    assert entradas[0] == "primera"

    liberar.set()
    for hilo in hilos:
        hilo.join(3.0)
        assert not hilo.is_alive()

    assert entradas == ["primera", "segunda"]


def test_bloqueo_existente_falla_al_vencer_el_tiempo(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path / "data")
    crear_bronze(settings.data_dir)
    bloqueo = settings.data_dir / ".pipeline.lock"

    bloqueo.mkdir(parents=True)
    pipeline = PipelineMedallon(settings, lock_timeout=0.0, lock_poll=0.01)

    with pytest.raises(ErrorConcurrencia):
        pipeline.ejecutar()
