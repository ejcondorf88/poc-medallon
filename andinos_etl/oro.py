"""CAPA ORO: modelo en estrella listo para Power BI (1 tabla de hechos + 5 dimensiones)."""
from __future__ import annotations

import pandas as pd
from openpyxl.utils import get_column_letter

from .base import CapaMedallon
from .config import CLAVE_DESCONOCIDA, DIAS_SEMANA, MESES, MESES_LARGOS
from .excepciones import ErrorIntegridad


class CapaOro(CapaMedallon):
    """Construye dimensiones con llaves sustitutas (enteros) y una tabla de hechos aditiva."""

    def ejecutar(self, plata: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
        print(
            "[ORO] Convierte datos validados en un modelo estrella para análisis: "
            "dimensiones con llaves sustitutas y un hecho con grano de línea de venta.",
            flush=True,
        )
        # Paso 1: dimensiones. Cada una recibe una llave sustituta (entero) estable y ordenada.
        print(
            "[ORO] Paso 1/4: crea las dimensiones ordenando IDs naturales para obtener "
            "llaves enteras reproducibles; cliente y promoción admiten miembro 0 "
            "solo cuando la referencia está vacía.",
            flush=True,
        )
        dim_tienda = self._dim_tienda(plata["tiendas"])
        dim_producto = self._dim_producto(plata["productos"])
        dim_cliente = self._dim_cliente(plata["clientes"])
        dim_promocion = self._dim_promocion(plata["promociones"])
        dim_fecha = self._dim_fecha(plata["ventas"])

        # Paso 2: tabla de hechos, que traduce los IDs de origen a las llaves sustitutas.
        print(
            "[ORO] Paso 2/4: convierte IDs en llaves dimensionales y calcula importes "
            "bruto, descuento, neto, costo y margen; una tienda o producto sin dimensión "
            "es error porque rompe relaciones analíticas.",
            flush=True,
        )
        fact_ventas = self._fact_ventas(plata["ventas"], dim_tienda, dim_producto, dim_cliente, dim_promocion)

        oro = {"dim_tienda": dim_tienda, "dim_producto": dim_producto, "dim_cliente": dim_cliente,
               "dim_promocion": dim_promocion, "dim_fecha": dim_fecha, "fact_ventas": fact_ventas}

        # Paso 3: persistir cada tabla como CSV individual.
        print(
            "[ORO] Paso 3/4: escribe CSV individuales para facilitar consumo e inspección.",
            flush=True,
        )
        for nombre, df in oro.items():
            self.almacen.escribir(df, self.settings.gold_dir / f"{nombre}.csv")
            print(f"[ORO]   {nombre}: {len(df):,} filas, {df.shape[1]} columnas.", flush=True)
            self.log.info("Oro    | %-13s | %6d filas | %2d columnas", nombre, len(df), df.shape[1])

        # Paso 4: un solo Excel con una pestaña por tabla (modelo en estrella).
        print(
            "[ORO] Paso 4/4: genera modelo_estrella.xlsx con una pestaña por tabla "
            "para inspección y carga en herramientas analíticas.",
            flush=True,
        )
        ruta_excel = self.settings.gold_dir / "modelo_estrella.xlsx"
        with pd.ExcelWriter(ruta_excel, engine="openpyxl") as writer:
            for nombre, df in oro.items():
                df_export = df.copy()
                for col in df_export.columns:
                    if df_export[col].dtype == object and len(df_export) > 0:
                        muestra = df_export[col].dropna()
                        if len(muestra) > 0 and hasattr(muestra.iloc[0], "year"):
                            df_export[col] = df_export[col].astype(str)
                df_export.to_excel(writer, sheet_name=nombre, index=False)
                hoja = writer.sheets[nombre]
                for indice, columna in enumerate(df_export.columns, start=1):
                    if not pd.api.types.is_datetime64_any_dtype(df_export[columna]):
                        continue
                    letra = get_column_letter(indice)
                    hoja.column_dimensions[letra].width = 20
                    for fila in range(2, len(df_export) + 2):
                        hoja.cell(row=fila, column=indice).number_format = "YYYY-MM-DD HH:MM:SS"
        self.log.info("Oro    | modelo_estrella.xlsx | %d pestañas generadas", len(oro))

        return oro

    # ---- utilidades --------------------------------------------------------
    @staticmethod
    def _con_llave(df: pd.DataFrame, natural: str, llave: str) -> pd.DataFrame:
        """Asigna llaves sustitutas 1..n ordenando por la llave natural (resultado reproducible)."""
        df = df.sort_values(natural).reset_index(drop=True)
        df.insert(0, llave, range(1, len(df) + 1))
        return df

    @staticmethod
    def _con_miembro_desconocido(df: pd.DataFrame, llave: str, valores: dict) -> pd.DataFrame:
        """Agrega la fila 0 ('Sin cliente', 'Sin promoción'): así el hecho nunca tiene llaves vacías."""
        fila = {c: pd.NA for c in df.columns} | {llave: CLAVE_DESCONOCIDA} | valores
        return pd.concat([pd.DataFrame([fila]), df], ignore_index=True)

    # ---- dimensiones -------------------------------------------------------
    def _dim_tienda(self, df: pd.DataFrame) -> pd.DataFrame:
        return self._con_llave(df.drop(columns="_fila_origen"), "tienda_id", "tienda_key")

    def _dim_producto(self, df: pd.DataFrame) -> pd.DataFrame:
        return self._con_llave(df.drop(columns="_fila_origen"), "producto_id", "producto_key")

    def _dim_cliente(self, df: pd.DataFrame) -> pd.DataFrame:
        dim = self._con_llave(df.drop(columns="_fila_origen"), "cliente_id", "cliente_key")
        dim = self._con_miembro_desconocido(dim, "cliente_key", {
            "cliente_id": "SIN_CLIENTE", "segmento": "Sin identificar",
            "ciudad_residencia": "Sin identificar", "canal_preferido": "Sin identificar"})
        dim["anio_nacimiento"] = dim["anio_nacimiento"].astype("Int64")
        return dim

    def _dim_promocion(self, df: pd.DataFrame) -> pd.DataFrame:
        dim = self._con_llave(df.drop(columns="_fila_origen"), "promocion_id", "promocion_key")
        return self._con_miembro_desconocido(dim, "promocion_key", {
            "promocion_id": "SIN_PROMO", "nombre_promocion": "Sin promoción",
            "tipo_promocion": "Sin promoción", "categoria_objetivo": "Sin promoción",
            "descuento_objetivo": 0.0})

    def _dim_fecha(self, ventas: pd.DataFrame) -> pd.DataFrame:
        """Calendario continuo por años completos: sin huecos, como exige la inteligencia de tiempo."""
        fechas = pd.to_datetime(ventas["fecha"])
        dias = pd.date_range(f"{fechas.min().year}-01-01", f"{fechas.max().year}-12-31", freq="D")
        print(
            f"[ORO] dim_fecha: calendario diario continuo {dias.min():%Y-%m-%d} a "
            f"{dias.max():%Y-%m-%d}; cubre años completos para evitar huecos en análisis temporales.",
            flush=True,
        )
        return pd.DataFrame({
            "fecha_key": dias.strftime("%Y%m%d").astype(int),
            "fecha": dias.date,
            "anio": dias.year,
            "trimestre": ["T" + str(q) for q in dias.quarter],
            "mes_num": dias.month,
            "mes": [MESES[m - 1] for m in dias.month],
            "mes_nombre": [MESES_LARGOS[m - 1] for m in dias.month],
            "anio_mes": dias.strftime("%Y-%m"),
            "dia": dias.day,
            "dia_semana_num": dias.dayofweek + 1,
            "dia_semana": [DIAS_SEMANA[d] for d in dias.dayofweek],
            "es_fin_de_semana": (dias.dayofweek >= 5).astype(int),
        })

    # ---- hechos ------------------------------------------------------------
    def _fact_ventas(self, ventas: pd.DataFrame, d_tienda: pd.DataFrame, d_producto: pd.DataFrame,
                     d_cliente: pd.DataFrame, d_promocion: pd.DataFrame) -> pd.DataFrame:
        f = ventas.copy()

        f = f.merge(d_tienda[["tienda_id", "tienda_key"]], on="tienda_id", how="left", validate="many_to_one")
        f = f.merge(d_producto[["producto_id", "producto_key"]], on="producto_id", how="left", validate="many_to_one")
        f = f.merge(d_cliente[["cliente_id", "cliente_key"]], on="cliente_id", how="left", validate="many_to_one")
        f = f.merge(d_promocion[["promocion_id", "promocion_key"]], on="promocion_id", how="left", validate="many_to_one")

        for natural, llave in (("tienda_id", "tienda_key"), ("producto_id", "producto_key")):
            huerfanos = f[llave].isna().sum()
            if huerfanos:
                raise ErrorIntegridad(f"{huerfanos} ventas con {natural} que no existe en su dimensión")

        for natural, llave in (("cliente_id", "cliente_key"), ("promocion_id", "promocion_key")):
            huerfanos = f[natural].notna() & f[llave].isna()
            if huerfanos.any():
                raise ErrorIntegridad(
                    f"{int(huerfanos.sum())} ventas con {natural} que no existe en su dimensión"
                )
            f[llave] = f[llave].fillna(CLAVE_DESCONOCIDA).astype(int)

        f["importe_bruto"] = f["cantidad"] * f["precio_unitario"]
        f["importe_descuento"] = f["importe_bruto"] * f["descuento_pct"]
        f["importe_neto"] = f["importe_bruto"] - f["importe_descuento"]
        f["costo_total"] = f["cantidad"] * f["costo_unitario"]
        f["margen"] = f["importe_neto"] - f["costo_total"]
        f["es_devolucion"] = (f["tipo_movimiento"] == "Devolución").astype(int)
        f["tiene_promocion"] = (f["promocion_key"] != CLAVE_DESCONOCIDA).astype(int)
        f["costo_imputado"] = f["costo_imputado"].astype(int)
        f["fecha_key"] = f["fecha_hora"].dt.strftime("%Y%m%d").astype(int)
        f["hora"] = f["fecha_hora"].dt.hour.astype(int)

        columnas = ["venta_id", "fecha_hora", "hora", "fecha_key", "tienda_key", "producto_key", "cliente_key",
                    "promocion_key", "cantidad", "precio_unitario", "descuento_pct", "costo_unitario",
                    "importe_bruto", "importe_descuento", "importe_neto", "costo_total", "margen",
                    "es_devolucion", "tiene_promocion", "costo_imputado", "medio_pago", "tipo_movimiento"]
        monetarias = ["importe_bruto", "importe_descuento", "importe_neto", "costo_total", "margen"]
        f[monetarias] = f[monetarias].round(4)
        return f[columnas].sort_values("venta_id").reset_index(drop=True)