"""Ejecutores independientes para las etapas Bronce, Plata, Oro y verificación."""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from . import artefactos
from .artefactos import copiar_bronce, escribir_plata, leer_plata
from .bronce import CapaBronce
from .calidad import CapaCalidad, ContextoCalidad
from .config import COLUMNAS_ESPERADAS, Settings
from .integridad import (
    CHECKSUMS_FILENAME,
    MANIFEST_FILENAME,
    construir_manifest,
    escribir_manifest,
    generar_checksums,
    leer_manifest,
    sha256_file,
    utc_now_iso,
    verificar_checksums,
)
from .oro import CapaOro
from .plata import CapaPlata

APPROVED_BUNDLE_NAME = "andinos-model-approved"


@dataclass(frozen=True)
class StageResult:
    stage: str
    status: str
    output_dir: Path
    files: tuple[Path, ...] = ()
    manifest: dict[str, Any] | None = None
    gate: dict[str, Any] | None = None
    errors: tuple[str, ...] = ()

    @property
    def ok(self) -> bool:
        return self.status in ("OK", "APPROVED")


@dataclass(frozen=True)
class ResultadoVerificacionArtifact:
    ok: bool
    errors: tuple[str, ...]


def _git_info() -> tuple[str, str]:
    version = os.environ.get("GITHUB_REF_NAME") or "local"
    commit = os.environ.get("GITHUB_SHA")
    if not commit:
        try:
            commit = subprocess.run(
                ["git", "rev-parse", "HEAD"], capture_output=True, check=True, text=True
            ).stdout.strip()
        except (OSError, subprocess.CalledProcessError):
            commit = "unknown"
    return version, commit


def _preparar_staging(destino: Path) -> Path:
    destino.parent.mkdir(parents=True, exist_ok=True)
    return Path(tempfile.mkdtemp(prefix=f".{destino.name}.staging-", dir=destino.parent))


def _promover(staging: Path, destino: Path) -> None:
    backup = destino.with_name(f".{destino.name}.backup-{os.getpid()}")
    if backup.exists():
        shutil.rmtree(backup)
    if destino.exists():
        os.replace(destino, backup)
    try:
        os.replace(staging, destino)
    except Exception:
        if backup.exists():
            os.replace(backup, destino)
        raise
    if backup.exists():
        shutil.rmtree(backup)


def _hashes_payload(raiz: Path, excluidos: set[str]) -> dict[str, str]:
    hashes: dict[str, str] = {}
    for ruta in sorted(raiz.rglob("*"), key=lambda item: item.relative_to(raiz).as_posix()):
        if not ruta.is_file():
            continue
        relativa = ruta.relative_to(raiz).as_posix()
        if relativa in excluidos:
            continue
        hashes[relativa] = sha256_file(ruta)
    return hashes


def _verificar_entrada(artifact_dir: Path) -> None:
    resultado = verificar_checksums(artifact_dir)
    if not resultado.ok:
        raise ValueError(f"Integridad inválida en {artifact_dir}: {'; '.join(resultado.errores)}")


def ejecutar_bronce(bronze_dir: Path, output_dir: Path) -> StageResult:
    print(
        "[ETAPA BRONCE] Construye un artefacto portable e inmutable. "
        "El manifiesto y los checksums permiten verificar origen e integridad después del transporte.",
        flush=True,
    )
    staging = _preparar_staging(output_dir)
    try:
        settings = Settings(data_dir=bronze_dir.parent, bronze_source=bronze_dir)
        datos_bronce = CapaBronce(settings).ejecutar()
        copiar_bronce(bronze_dir, staging)
        for tabla, df in datos_bronce.items():
            df[COLUMNAS_ESPERADAS[tabla]].to_csv(
                staging / f"{tabla}.csv",
                index=False,
                encoding="utf-8-sig",
                lineterminator="\n",
            )
        version, commit = _git_info()
        created_at = utc_now_iso()
        metadata = {
            "schema_version": 1,
            "stage": "bronce",
            "created_at_utc": created_at,
            "git_version": version,
            "git_commit": commit,
            "source": bronze_dir.name,
            "immutable": True,
            "rows_by_table": {tabla: len(df) for tabla, df in datos_bronce.items()},
        }
        (staging / "stage.json").write_text(
            json.dumps(metadata, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
        )
        manifest = construir_manifest(
            "bronce",
            _hashes_payload(staging, {MANIFEST_FILENAME, CHECKSUMS_FILENAME}),
            git_version=version,
            git_commit=commit,
            created_at_utc=created_at,
            extra={"source": bronze_dir.name, "immutable": True},
        )
        escribir_manifest(manifest, staging / MANIFEST_FILENAME)
        generar_checksums(staging)
        _promover(staging, output_dir)
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    return StageResult("bronce", "OK", output_dir, tuple(sorted(output_dir.rglob("*"))), manifest)


def ejecutar_plata(bronze_dir: Path, output_dir: Path) -> StageResult:
    print(
        "[ETAPA PLATA] Verifica primero los checksums de Bronce; solo procesa una entrada "
        "íntegra para que la limpieza y su bitácora se puedan reproducir.",
        flush=True,
    )
    staging = _preparar_staging(output_dir)
    try:
        _verificar_entrada(bronze_dir)
        datos_bronce = artefactos.leer_bronze(bronze_dir)
        staging_silver = staging / "_silver"
        staging_quality = staging / "_quality"
        settings = Settings(
            data_dir=bronze_dir.parent,
            silver_source=staging_silver,
            quality_source=staging_quality,
        )
        resultados = CapaPlata(settings).ejecutar(datos_bronce)
        escribir_plata(staging, resultados)
        shutil.rmtree(staging_silver, ignore_errors=True)
        shutil.rmtree(staging_quality, ignore_errors=True)
        version, commit = _git_info()
        manifest = construir_manifest(
            "plata",
            _hashes_payload(staging, {MANIFEST_FILENAME, CHECKSUMS_FILENAME}),
            git_version=version,
            git_commit=commit,
            created_at_utc=utc_now_iso(),
            rows_processed=sum(len(resultado.limpio) for resultado in resultados.values()),
        )
        escribir_manifest(manifest, staging / MANIFEST_FILENAME)
        generar_checksums(staging)
        _promover(staging, output_dir)
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    return StageResult("plata", "OK", output_dir, tuple(sorted(output_dir.rglob("*"))), manifest)


def _crear_gate(reporte: pd.DataFrame) -> dict[str, Any]:
    controles = [
        {
            "control": str(fila["control"]),
            "descripcion": str(fila["descripcion"]),
            "esperado": str(fila["esperado"]),
            "obtenido": str(fila["obtenido"]),
            "estado": str(fila["estado"]),
        }
        for fila in reporte.to_dict("records")
    ]
    fallas = [control for control in controles if control["estado"] == "FALLO"]
    advertencias = [control for control in controles if control["estado"] == "ADVERTENCIA"]
    return {
        "schema_version": 1,
        "status": "REJECTED" if fallas else "APPROVED",
        "failures": len(fallas),
        "warnings": len(advertencias),
        "controls": controles,
    }


def _escribir_gate(ruta: Path, gate: dict[str, Any]) -> None:
    ruta.write_text(json.dumps(gate, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def ejecutar_oro(bronze_dir: Path, plata_dir: Path, output_dir: Path) -> StageResult:
    print(
        "[ETAPA ORO] Verifica checksums de Bronce y Plata antes de modelar. "
        "El gate reúne controles; únicamente APPROVED produce un bundle consumible.",
        flush=True,
    )
    staging = _preparar_staging(output_dir)
    try:
        _verificar_entrada(bronze_dir)
        _verificar_entrada(plata_dir)
        datos_bronce = artefactos.leer_bronze(bronze_dir)
        resultados_plata = leer_plata(plata_dir)
        staging_gold = staging / "_gold"
        staging_quality = staging / "_quality"
        settings = Settings(
            data_dir=bronze_dir.parent,
            gold_source=staging_gold,
            quality_source=staging_quality,
        )
        datos_oro = CapaOro(settings).ejecutar({tabla: resultado.limpio for tabla, resultado in resultados_plata.items()})
        contexto = ContextoCalidad(datos_bronce, resultados_plata, datos_oro)
        reporte = CapaCalidad(settings).ejecutar(contexto)
        gate = _crear_gate(reporte)
        if gate["status"] == "REJECTED":
            print(
                f"[ETAPA ORO] REJECTED: {gate['failures']} fallo(s), "
                f"{gate['warnings']} advertencia(s). No se crea el bundle aprobado; "
                "se conserva diagnóstico para corregir la causa.",
                flush=True,
            )
            diagnostico = staging / "diagnostico"
            (diagnostico / "calidad").mkdir(parents=True, exist_ok=True)
            shutil.copy2(staging_quality / "reporte_calidad.csv", diagnostico / "calidad" / "reporte_calidad.csv")
            reglas = plata_dir / "reglas_limpieza.csv"
            if reglas.is_file():
                shutil.copy2(reglas, diagnostico / "calidad" / "reglas_limpieza.csv")
            _escribir_gate(diagnostico / "gate.json", gate)
            shutil.rmtree(staging_gold, ignore_errors=True)
            shutil.rmtree(staging_quality, ignore_errors=True)
            _promover(staging, output_dir)
            return StageResult("oro", "REJECTED", output_dir, tuple(sorted(output_dir.rglob("*"))), gate=gate)
        bundle = staging / APPROVED_BUNDLE_NAME
        modelo = bundle / "modelo_estrella"
        calidad = bundle / "calidad"
        modelo.mkdir(parents=True)
        calidad.mkdir(parents=True)
        for nombre, df in datos_oro.items():
            df.to_csv(modelo / f"{nombre}.csv", index=False, encoding="utf-8-sig", date_format="%Y-%m-%dT%H:%M:%S", lineterminator="\n")
        shutil.copy2(staging_gold / "modelo_estrella.xlsx", bundle / "modelo_estrella.xlsx")
        shutil.copy2(staging_quality / "reporte_calidad.csv", calidad / "reporte_calidad.csv")
        shutil.copy2(plata_dir / "reglas_limpieza.csv", calidad / "reglas_limpieza.csv")
        _escribir_gate(bundle / "gate.json", gate)
        hechos = datos_oro["fact_ventas"]
        ventas_netas = round(float(hechos["importe_neto"].sum()), 2)
        costo_total = round(float(hechos["costo_total"].sum()), 2)
        version, commit = _git_info()
        manifest = construir_manifest(
            "oro",
            _hashes_payload(bundle, {MANIFEST_FILENAME, CHECKSUMS_FILENAME}),
            git_version=version,
            git_commit=commit,
            created_at_utc=utc_now_iso(),
            rows_processed=len(datos_bronce["ventas"]),
            ventas_netas=ventas_netas,
            costo_total=costo_total,
            margen=round(ventas_netas - costo_total, 2),
            quality_gate=gate,
            extra={
                "fact_rows": len(hechos),
                "rows_by_table": {nombre: len(df) for nombre, df in datos_oro.items()},
            },
        )
        escribir_manifest(manifest, bundle / MANIFEST_FILENAME)
        generar_checksums(bundle)
        shutil.rmtree(staging_gold, ignore_errors=True)
        shutil.rmtree(staging_quality, ignore_errors=True)
        _promover(staging, output_dir)
        approved = output_dir / APPROVED_BUNDLE_NAME
        print(
            f"[ETAPA ORO] APPROVED: {gate['warnings']} advertencia(s), 0 fallos. "
            "Bundle con modelo, reporte de calidad, reglas de limpieza, manifiesto y checksums.",
            flush=True,
        )
    except Exception:
        shutil.rmtree(staging, ignore_errors=True)
        raise
    return StageResult("oro", "APPROVED", approved, tuple(sorted(approved.rglob("*"))), manifest, gate)


def _conteos_csv(modelo_dir: Path) -> dict[str, int]:
    return {
        ruta.stem: int(len(pd.read_csv(ruta, encoding="utf-8-sig")))
        for ruta in sorted(modelo_dir.glob("*.csv"))
    }


def _conteos_xlsx(ruta: Path) -> dict[str, int]:
    conteos: dict[str, int] = {}
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        for hoja in libro.worksheets:
            filas = sum(
                1
                for fila in hoja.iter_rows(min_row=2, values_only=True)
                if any(valor is not None for valor in fila)
            )
            conteos[hoja.title] = filas
    finally:
        libro.close()
    return conteos


def _verificar_conteos(artifact_dir: Path, manifest: dict[str, Any]) -> list[str]:
    esperados = manifest.get("rows_by_table")
    if not isinstance(esperados, dict) or not esperados:
        return ["El manifest no registra rows_by_table"]
    modelo_dir = artifact_dir / "modelo_estrella"
    xlsx = artifact_dir / "modelo_estrella.xlsx"
    if not modelo_dir.is_dir():
        return ["Falta la carpeta: modelo_estrella"]
    if not xlsx.is_file():
        return ["Falta archivo: modelo_estrella.xlsx"]
    try:
        csvs = _conteos_csv(modelo_dir)
        hojas = _conteos_xlsx(xlsx)
    except Exception as error:
        return [f"No se pudieron leer los conteos del bundle: {type(error).__name__}: {error}"]
    errores: list[str] = []
    for tabla in sorted(set(esperados) | set(csvs) | set(hojas)):
        esperado = esperados.get(tabla)
        filas_csv = csvs.get(tabla)
        filas_hoja = hojas.get(tabla)
        if esperado is None:
            errores.append(f"{tabla}: presente en el bundle pero no en rows_by_table")
        elif filas_csv is None:
            errores.append(f"{tabla}: falta el CSV modelo_estrella/{tabla}.csv")
        elif filas_hoja is None:
            errores.append(f"{tabla}: falta la hoja {tabla} en modelo_estrella.xlsx")
        elif not esperado == filas_csv == filas_hoja:
            errores.append(
                f"Conteos distintos en {tabla}: manifest {esperado} | CSV {filas_csv} | Excel {filas_hoja}"
            )
    return errores


def verificar_artifact(artifact_dir: Path) -> StageResult:
    print(
        "[VERIFICAR] Comprueba hashes para detectar cambios, valida manifest y gate, "
        "y compara conteos de filas entre CSV, Excel y manifiesto para detectar artefactos incompletos.",
        flush=True,
    )
    resultado_checksums = verificar_checksums(artifact_dir)
    errores = list(resultado_checksums.errores)
    manifest_path = artifact_dir / MANIFEST_FILENAME
    manifest: dict[str, Any] | None = None
    if not manifest_path.is_file():
        errores.append("Falta archivo: manifest.json")
    else:
        try:
            manifest = leer_manifest(manifest_path)
        except (json.JSONDecodeError, OSError, ValueError) as error:
            errores.append(f"Manifest inválido: {error}")
    stage = manifest.get("stage") if manifest else None
    gate_path = artifact_dir / "gate.json"
    gate: dict[str, Any] | None = None
    if gate_path.is_file():
        try:
            gate = json.loads(gate_path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError) as error:
            errores.append(f"Gate inválido: {error}")
    if stage == "oro":
        if not gate_path.is_file():
            errores.append("Falta archivo: gate.json")
        elif gate and gate.get("status") != "APPROVED":
            errores.append(f"Quality gate no aprobado: {gate.get('status')}")
        if manifest is not None:
            errores.extend(_verificar_conteos(artifact_dir, manifest))
    if errores:
        print(f"[VERIFICAR] ERROR: {len(errores)} problema(s) detectado(s).", flush=True)
        for error in errores:
            print(f"[VERIFICAR]   - {error}", flush=True)
    else:
        print("[VERIFICAR] OK: integridad, gate y conteos verificados.", flush=True)
    return StageResult(
        "verificar",
        "OK" if not errores else "ERROR",
        artifact_dir,
        manifest=manifest,
        gate=gate,
        errors=tuple(errores),
    )
